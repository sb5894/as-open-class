# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl==3.1.5"]
# ///
"""Read-only XLSX import; deterministic site JSON, selected PDFs and review reports.

Run from the project root: uv run scripts/convert_data.py [--workbook path] [--check]
Room and PDF assignments are reviewed rules in data_rules.py, never inferred from
a grade number or an unrelated PDF's contents. No workbook/PDF is rewritten.
"""
import argparse
from datetime import date, datetime
from hashlib import sha256
import json
from pathlib import Path
import re
import sys
import unittest

import openpyxl

from data_rules import (LENSES, MULTICULTURAL, PDF_RULES, PRIVATE_CLASSES, SPECIAL_LESSONS,
                        WORKBOOK_MISSING_CLASSES, WORKBOOK_SPECIAL_CLASSES, spatial_data)

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "materials/original/curriculum/2026 공개보고회 준비 수업정보(261015-2113).xlsx"
SUPPLEMENT_WORKBOOK = ROOT / "materials/original/curriculum/1020-additions/2026 공개보고회 준비(공개수업, 예산).xlsx"
FLOOR_PLAN = ROOT / "materials/original/floor-plans/안성초 도면.pdf"
GRADE_SHEET = "공개수업(학년)"
CLASS_SHEET = "공개수업(학급, 전담, 특별학급)"
TEXT_FIELDS = ("gradeTheme", "topic", "designFocus", "commonObservation", "context", "observation", "support")
FIELD_NAMES = {"gradeTheme": "학년군 주제", "topic": "수업 주제", "designFocus": "설계 포인트",
               "commonObservation": "공통 참관 포인트", "context": "학급·수업 맥락", "observation": "학급 참관 포인트",
               "support": "지원인력·환경", "lenses": "중점 렌즈", "lessonPlanId": "지도안 PDF"}


def text(value):
    if value is None:
        return None
    result = str(value).replace("\r\n", "\n").strip()
    return result or None


def normalize_class(value):
    """openpyxl resolves Excel's date system (1900/1904) before this step."""
    if isinstance(value, (datetime, date)):
        label = f"{value.month}-{value.day}"
    elif isinstance(value, str):
        match = re.fullmatch(r"\s*([1-6])\s*[-－]\s*(\d{1,2})(?:\s*\([^)]*\))?\s*", value)
        if not match:
            raise ValueError(f"인식할 수 없는 학급명: {value!r}")
        label = f"{int(match[1])}-{int(match[2])}"
    else:
        raise ValueError(f"학급명은 날짜형 또는 학년-반 문자열이어야 합니다: {value!r}")
    grade, number = map(int, label.split("-"))
    if grade not in range(1, 7) or number not in range(1, 31):
        raise ValueError(f"학급명 범위 오류: {label}")
    return label


def merged_value(sheet, row, col):
    cell = sheet.cell(row, col)
    if cell.value is not None:
        return cell_value(sheet, row, col)
    for area in sheet.merged_cells.ranges:
        if cell.coordinate in area:
            return cell_value(sheet, area.min_row, area.min_col)
    return None


def source_ref(sheet, row):
    return {"sheet": sheet.title, "row": row}


def cell_value(sheet, row, col):
    cell = sheet.cell(row, col)
    if cell.data_type == "f":
        raise ValueError(f"수식 셀은 직접 확인 후 값으로 제공해 주세요: {sheet.title}!{cell.coordinate}")
    return cell.value


def class_targets(value):
    if isinstance(value, (datetime, date)):
        return [normalize_class(value)]
    value = text(value) or ""
    # Remove parenthetical exceptions with their own class label, preserving
    # targets before them: '6-2, 6-4, 6-5 (6-1 영어 별도)'.
    value = re.sub(r"\([^)]*[1-6]\s*-\s*\d{1,2}[^)]*(?:별도|비공개)[^)]*\)", "", value)
    targets = []
    for part in re.split(r"[,;\n]", value):
        if "별도" in part or "비공개" in part:
            continue
        part = re.sub(r"\([^)]*\)", "", part)
        targets.extend(normalize_class(label) for label in re.findall(r"[1-6]\s*-\s*\d{1,2}", part))
    return targets


def read_grade_topics(sheet):
    candidates = {}
    for row in range(1, sheet.max_row + 1):
        label = text(merged_value(sheet, row, 1))
        if not label or not re.fullmatch(r"[1-6]학년", label):
            continue
        grade = int(label[0])
        item = {"gradeTheme": text(merged_value(sheet, row, 2)),
                "topic": text(cell_value(sheet, row, 3)),
                "designFocus": text(cell_value(sheet, row, 4)),
                "commonObservation": text(cell_value(sheet, row, 5)),
                "source": source_ref(sheet, row), "targets": class_targets(cell_value(sheet, row, 6))}
        candidates.setdefault(grade, []).append(item)
    if set(candidates) != set(range(1, 7)):
        raise ValueError("학년 시트에서 1~6학년 행을 모두 찾을 수 없습니다. 양식을 확인해 주세요.")
    # Explicit targeting is mandatory for multiple topics, never assume the first.
    for grade, items in candidates.items():
        targets = [target for item in items for target in item["targets"]]
        if len(items) > 1 and (any(not item["targets"] for item in items) or len(set(targets)) != len(targets)):
            raise ValueError(f"{grade}학년 여러 주제의 학급 구분이 없거나 중복됩니다.")
    return candidates


def read_class_rows(sheet, special_classes=WORKBOOK_SPECIAL_CLASSES):
    rows = {}
    section = None
    for row in range(1, sheet.max_row + 1):
        raw = cell_value(sheet, row, 1)
        label = text(raw)
        if not label:
            continue
        if re.match(r"^[1-6]학년(?:\s|$)", label):
            section = "regular"
            continue
        if label == "일반 교과수업 적용 사례":
            section = "special"
            continue
        if label == "다문화특별학급 적용 사례":
            section = "multicultural"
            continue
        if section is None or label in {"학급", "렌즈 1"} or label.startswith("※"):
            continue
        if section == "multicultural":
            match = re.fullmatch(r"다문화특별학급\s*([123])반\((빨강|초록|파랑)\)", label)
            if not match:
                raise ValueError(f"다문화특별학급 표기를 확인해 주세요: A{row} {label!r}")
            info = MULTICULTURAL[match[1]]
            if match[2] not in info["name"]:
                raise ValueError(f"다문화 반 번호와 색상 불일치: A{row}")
            class_id = info["id"]
        else:
            class_id = normalize_class(raw)
            if (class_id in special_classes) != (section == "special"):
                raise ValueError(f"별도 수업 구분이 결정 기록과 다릅니다: {class_id} (행 {row})")
        if class_id in rows:
            raise ValueError(f"학급 행이 중복됩니다: {class_id}")
        lenses = [text(cell_value(sheet, row, col)) for col in (4, 5, 6)]
        lenses = [lens for lens in lenses if lens is not None]
        if any(lens not in LENSES for lens in lenses) or len(set(lenses)) != len(lenses):
            raise ValueError(f"중점 렌즈 오기 또는 중복: {class_id} {lenses}")
        rows[class_id] = {"context": text(cell_value(sheet, row, 2)),
                          "observation": text(cell_value(sheet, row, 3)), "lenses": lenses,
                          "support": text(cell_value(sheet, row, 7)), "source": source_ref(sheet, row)}
    return rows


def choose_topic(class_id, grade, candidates):
    items = candidates[grade]
    if len(items) == 1:
        return items[0]
    matches = [item for item in items if class_id in item["targets"]]
    if len(matches) != 1:
        raise ValueError(f"수업 주제를 하나로 연결할 수 없습니다: {class_id}")
    return matches[0]


def file_ref(path):
    return {"path": path.relative_to(ROOT).as_posix(), "sha256": sha256(path.read_bytes()).hexdigest()}


def validate_catalog(data):
    def unique(items, name):
        ids = [item["id"] for item in items]
        if len(ids) != len(set(ids)):
            raise ValueError(f"{name} ID 중복")
        return {item["id"]: item for item in items}

    buildings = unique(data["buildings"], "건물")
    spaces = unique(data["spaces"], "공간")
    classes = unique(data["classes"], "학급")
    plans = unique(data["lessonPlans"], "지도안")
    unique(data["lessons"], "수업")
    for space in spaces.values():
        building = buildings.get(space["buildingId"])
        if building is None or space["floor"] not in building["floors"]:
            raise ValueError(f"공간의 건물·층 오류: {space['id']}")
    for school_class in classes.values():
        if not school_class["homeSpaceIds"] or any(room not in spaces for room in school_class["homeSpaceIds"]):
            raise ValueError(f"원래 교실 연결 오류: {school_class['id']}")
    lesson_classes = []
    for lesson in data["lessons"]:
        class_id = lesson["classId"]
        school_class = classes.get(class_id)
        if school_class is None or not school_class["isPublic"]:
            raise ValueError(f"공개 대상이 아닌 학급의 수업: {class_id}")
        lesson_classes.append(class_id)
        if not lesson["spaceIds"] or any(room not in spaces or spaces[room]["kind"] == "staff" for room in lesson["spaceIds"]):
            raise ValueError(f"실제 수업 장소 오류: {class_id}")
        expected_rooms = ([SPECIAL_LESSONS[class_id]["roomId"]] if class_id in SPECIAL_LESSONS else school_class["homeSpaceIds"])
        if lesson["spaceIds"] != expected_rooms:
            raise ValueError(f"결정 기록과 실제 수업 장소가 다릅니다: {class_id}")
        plan_id = lesson["lessonPlanId"]
        if plan_id is not None and (plan_id not in plans or class_id not in PDF_RULES[plan_id]["classes"]):
            raise ValueError(f"잘못된 지도안 연결: {class_id}")
        if class_id in SPECIAL_LESSONS and (lesson["source"]["gradeRow"] is not None
                                            or (plan_id is not None and plan_id != f"special-{class_id}")):
            raise ValueError(f"별도 수업에 공통 자료를 연결했습니다: {class_id}")
        if lesson["locationStatus"] != "confirmed":
            raise ValueError(f"장소 확인 상태 오류: {class_id}")
    expected = {key for key, item in classes.items() if item["isPublic"]}
    if set(lesson_classes) != expected or len(lesson_classes) != len(expected):
        raise ValueError("공개수업 누락 또는 한 학급에 여러 수업 연결")


def convert(workbook_path, supplement_workbook=SUPPLEMENT_WORKBOOK):
    workbook_path = Path(workbook_path).resolve()
    if not workbook_path.is_relative_to(ROOT) or workbook_path.suffix.lower() != ".xlsx":
        raise ValueError("프로젝트 안의 .xlsx 파일 경로를 지정해 주세요.")
    workbook = openpyxl.load_workbook(workbook_path, data_only=False)
    try:
        candidates = read_grade_topics(workbook[GRADE_SHEET])
        class_rows = read_class_rows(workbook[CLASS_SHEET])
    finally:
        workbook.close()
    spatial = spatial_data()
    expected = {item["id"] for item in spatial["classes"] if item["id"] not in WORKBOOK_MISSING_CLASSES}
    if set(class_rows) != expected:
        raise ValueError(f"학급 목록이 공간 규칙과 다릅니다. 누락={sorted(expected-set(class_rows))}, 추가={sorted(set(class_rows)-expected)}")
    lesson_plans, pdf_files, plan_by_class = [], {}, {}
    missing_files = []
    for plan_id, rule in PDF_RULES.items():
        if not any(item['isPublic'] and item['id'] in rule['classes'] for item in spatial['classes']):
            continue
        source = ROOT / "materials/original/lesson-plans" / rule["file"]
        if not source.exists():
            missing_files.append(rule["file"])
            continue
        public_path = f"public/pdfs/{plan_id}.pdf"
        pdf_files[public_path] = source.read_bytes()
        lesson_plans.append({"id": plan_id, "url": f"/pdfs/{plan_id}.pdf", "source": file_ref(source)})
        for class_id in rule["classes"]:
            if class_id in plan_by_class:
                raise ValueError(f"PDF 연결 규칙 중복: {class_id}")
            plan_by_class[class_id] = plan_id
    lessons = []
    for school_class in spatial["classes"]:
        if not school_class["isPublic"]:
            continue
        class_id = school_class["id"]
        row = class_rows.get(class_id)
        special = SPECIAL_LESSONS.get(class_id)
        grade_topic = None
        if school_class["kind"] == "homeroom" and special is None:
            grade_topic = choose_topic(class_id, school_class["grade"], candidates)
        lessons.append({
            "id": "lesson-" + class_id, "classId": class_id,
            "kind": "special" if special else school_class["kind"],
            "subject": special["subject"] if special else None,
            "spaceIds": [special["roomId"]] if special else school_class["homeSpaceIds"],
            "locationStatus": "confirmed",
            **{key: grade_topic[key] if grade_topic else None for key in TEXT_FIELDS[:4]},
            **{key: row[key] if row else None for key in ("context", "observation", "support")},
            "lenses": row["lenses"] if row else [],
            "lessonPlanId": plan_by_class.get(class_id),
            "source": {"classRow": row["source"] if row else None, "gradeRow": grade_topic["source"] if grade_topic else None},
        })
    sources = {"workbook": file_ref(workbook_path), "floorPlan": file_ref(FLOOR_PLAN), "decisions": "docs/결정 기록.md"}
    if supplement_workbook is not None:
        supplement_workbook = Path(supplement_workbook).resolve()
        if not supplement_workbook.is_relative_to(ROOT) or supplement_workbook.suffix.lower() != ".xlsx":
            raise ValueError("추가 자료는 프로젝트 안의 .xlsx 파일이어야 합니다.")
        workbook = openpyxl.load_workbook(supplement_workbook, data_only=False)
        try:
            supplement_rows = read_class_rows(workbook[CLASS_SHEET], special_classes=set(SPECIAL_LESSONS))
            supplement_topics = read_grade_topics(workbook[GRADE_SHEET])
        finally:
            workbook.close()
        expected_supplement = {item["id"] for item in spatial["classes"] if item["isPublic"] and item["id"] != "kindergarten-joint"}
        if set(supplement_rows) != expected_supplement:
            raise ValueError(f"추가 자료의 공개 학급 목록이 다릅니다. 누락={sorted(expected_supplement-set(supplement_rows))}, 추가={sorted(set(supplement_rows)-expected_supplement)}")
        sources["supplementalWorkbook"] = file_ref(supplement_workbook)
        for lesson in lessons:
            if lesson["kind"] == "homeroom":
                grade = next(item["grade"] for item in spatial["classes"] if item["id"] == lesson["classId"])
                topic = choose_topic(lesson["classId"], grade, supplement_topics)
                changed = []
                for field in ("topic", "designFocus", "commonObservation"):
                    if topic[field] and topic[field] != lesson[field]:
                        lesson[field] = topic[field]
                        changed.append(field)
                if changed:
                    lesson["source"]["supplementalGradeRow"] = {**topic["source"], "fields": changed}
            row = supplement_rows.get(lesson["classId"])
            if row is None:
                continue
            changed = []
            for field in ("context", "observation", "support", "lenses"):
                # Empty new cells never erase existing content; only explicit new text wins.
                if row[field] and row[field] != lesson[field]:
                    lesson[field] = row[field]
                    changed.append(field)
            if changed:
                lesson["source"]["supplementalClassRow"] = {**row["source"], "fields": changed}
    data = {"schemaVersion": 1, "sources": sources,
            "event": {"name": None, "date": "2026-11-11", "time": None, "address": None},
            **spatial, "lessons": lessons, "lessonPlans": lesson_plans,
            "excludedClasses": [{"classId": key, "reason": reason} for key, reason in PRIVATE_CLASSES.items()]}
    validate_catalog(data)
    return data, pdf_files, missing_files


def location_label(data, ids):
    buildings = {item["id"]: item["name"] for item in data["buildings"]}
    spaces = {item["id"]: item for item in data["spaces"]}
    return " / ".join(f"{buildings[spaces[key]['buildingId']]} {spaces[key]['floor']}층 {spaces[key]['name']}" for key in ids)


def md(value):
    return (str(value) if value is not None else "미제공").replace("|", "\\|").replace("\n", "<br>")


def reports(data, missing_files):
    classes = {item["id"]: item for item in data["classes"]}
    plans = {item["id"]: item for item in data["lessonPlans"]}
    review = ["# 수업·장소·지도안 연결표", "", "자동 생성 문서입니다. `npm.cmd run data:convert`로 갱신합니다.", "",
              f"기준 엑셀: `{data['sources']['workbook']['path']}`", "",
              f"공개수업 {len(data['lessons'])}개. 학급과 실제 장소를 따로 관리합니다. 유치원 공동 수업은 본관 1층 유치원으로 안내합니다.", "",
              "| 수업 ID | 학급·수업 | 수업 주제 | 실제 참관 장소 | 장소 상태 | 지도안 | 엑셀 근거 |", "|---|---|---|---|---|---|---|"]
    missing = ["# 누락 자료와 확인할 사항", "", "자동 생성 문서입니다. 데이터의 미제공 텍스트·PDF는 `null`, 작성된 렌즈가 없으면 `[]`입니다.", "",
               "화면에서는 `src/components/placeholders.ts`를 사용합니다. 아래 ‘미제공’은 검토표용 표현입니다.", "",
               "- 행사 세부 시간. 정식 행사명·주소는 사용자 확정 설정 `src/data/event.ts`와 결정 기록을 따름.", "- 유치원 공동 수업은 본관 1층 유치원에서 공개함. 텍스트·PDF 제공 여부는 아래 표와 수업 연결표를 따르며 원본과 공간 기록은 보존.",
               "- 주차 진입 경로와 본관 중앙 출입구 현장 확인(가안).", "",
               "## 수업별 미제공 항목", "", "설계·공통 참관 포인트는 일반 학년 공통 수업에만 적용합니다. 별도·다문화·유치원에는 억지로 요구하지 않습니다.", "",
               "| 학급·수업 | 미제공 항목 |", "|---|---|"]
    for lesson in data["lessons"]:
        school_class = classes[lesson["classId"]]
        source = lesson["source"]
        refs = [f"{ref['sheet']}!행 {ref['row']}" for ref in (source["classRow"], source["gradeRow"]) if ref]
        title = school_class["name"] + (f" {lesson['subject']}" if lesson["subject"] else "")
        plan = plans.get(lesson["lessonPlanId"])
        pdf = f"[{plan['id']}](../public{plan['url']})" if plan else "미제공(공통 대체 없음)"
        review.append("| " + " | ".join(md(x) for x in [lesson["id"], title, lesson["topic"], location_label(data, lesson["spaceIds"]),
                                                        "가안" if lesson["locationStatus"] == "provisional" else "결정 기록 반영", pdf, "; ".join(refs) or "결정 기록"]) + " |")
        relevant = ["topic", "context", "observation", "support", "lenses", "lessonPlanId"]
        if lesson["kind"] == "homeroom":
            relevant = list(TEXT_FIELDS) + ["lenses", "lessonPlanId"]
        absent = [FIELD_NAMES[key] for key in relevant if lesson[key] is None or lesson[key] == []]
        if absent:
            missing.append(f"| {md(title)} | {', '.join(absent)} |")
    review += ["", "## 공간 연결 유의사항", "", "- 5-1 도덕 원래 교실: 본관 3층 → 실제 수업: 본관 2층 1과학실.",
               "- 4-5 원래 교실: 후관 3층 → 실제 수업: 본관 2층 2과학실.",
               "- 6-1 원래 교실: 본관 4층 → 실제 수업: 같은 층 영어실.", "- 1-3 수석교사 수업: 후관 2층 1-3 교실.",
               "- 한국어교실 북쪽 공간: 왼쪽부터 빨강반 / 강사용 1 / 강사용 2 / 초록반. 파랑반은 교장실 옆 남쪽 공간.",
               "- 3-1 체육과 강사용 교실 2개는 공개수업 및 희망 수업 대상에서 제외.",
               "- 공간 순서는 도면에서 추출한 상대 순서이며 실측 좌표가 아님. 계단·승강기·화장실과 전체 SVG는 6단계 범위.", ""]
    missing += ["", "## 원문을 보존하며 확인할 내용", ""]
    by_class = {lesson["classId"]: lesson for lesson in data["lessons"]}
    issues = []
    if by_class["multicultural-red"]["observation"] == "한국어 표현에 어려움이 있는 다문화 학생들이":
        issues.append("- 빨강반 참관 포인트는 ‘한국어 표현에 어려움이 있는 다문화 학생들이’에서 끝남. 작성 완료 여부 확인.")
    common_points = {}
    for lesson in data["lessons"]:
        if lesson["commonObservation"]:
            common_points.setdefault(lesson["commonObservation"], {})[json.dumps(lesson["source"]["gradeRow"], sort_keys=True)] = lesson["source"]["gradeRow"]
    for lesson in data["lessons"]:
        matches = common_points.get(lesson["observation"], {})
        if matches:
            refs = ", ".join(f"{ref['sheet']}!E{ref['row']}" for ref in matches.values())
            row = lesson["source"]["classRow"]
            issues.append(f"- {classes[lesson['classId']]['name']} 참관 포인트({row['sheet']}!C{row['row']})가 학년 공통 참관 포인트({refs})와 같음. 작성자 확인 필요. 원문 자동 교정 없음.")
    missing += issues or ["- 자동 비교로 확인된 원문 검토 항목 없음. 내용의 정확성은 작성자가 확인해야 함."]
    missing += ["", "## 원본 PDF와 결정 기록의 대조 — 작성자 확인 필요", "",
                "아래는 2026-10-16에 제공된 PDF를 직접 확인한 기록이다. 새 PDF가 오면 다시 확인한다. 데이터/PDF 연결 규칙은 결정 기록을 유지하며 원문을 자동 교체하지 않는다.", "",
                "- 4학년 PDF 1쪽의 4-3 고유 문구: ‘한국어 설명과 러시아어 표현의 역할을 나누어 서로의 설명을 보완한다’. 엑셀 C32와의 차이를 작성자가 확인해야 하며, PDF 문구로 엑셀을 자동 덮어쓰지 않음.",
                "- 3학년 PDF 1쪽에는 3-1의 언어·문화 경험 활용 계획이 있음. 현재 사용자 결정으로 3-1 체육은 비공개이며 원본 PDF는 변경하지 않음.",
                "- 4학년 PDF 1쪽에는 4-5의 표현 지원 계획이 있음. 결정 기록/엑셀에서는 4-5 과학이 별도 수업임.",
                "- 5학년 PDF 2쪽에는 5-1의 학급 설명이 있음. 현재 5-1 도덕은 본관 2층 1과학실의 별도 공개수업이며 기존 5학년 PDF를 연결하지 않음.",
                ("- 유치원 지도안이 제공되어 해당 수업의 PDF로 연결함. PDF 연결만 반영하며 수업 설명 본문은 기존 엑셀 원문을 유지함."
                 if by_class["kindergarten-joint"]["lessonPlanId"] else
                 "- 유치원 지도안이 제공되지 않아 PDF 연결 없이 미제공 표시를 유지함.")]
    if missing_files:
        missing += ["", "## 기존 연결 규칙의 파일 누락", ""] + [f"- {name}" for name in missing_files]
    missing.append("")
    return {"docs/수업 연결표.md": "\n".join(review).encode("utf-8"),
            "docs/누락 자료.md": "\n".join(missing).encode("utf-8")}


def output_files(data, pdf_files, missing_files):
    classes = {item["id"]: item for item in data["classes"]}
    registration_lessons = [{"id": lesson["id"],
                             "label": classes[lesson["classId"]]["name"] + (f" {lesson['subject']}" if lesson["subject"] else "")
                                      + " · " + location_label(data, lesson["spaceIds"])} for lesson in data["lessons"]]
    return {"src/data/catalog.generated.json": (json.dumps(data, ensure_ascii=False, indent=2) + "\n").encode("utf-8"),
            "functions/src/data/registration-lessons.generated.json": (json.dumps(registration_lessons, ensure_ascii=False, indent=2) + "\n").encode("utf-8"),
            **pdf_files, **reports(data, missing_files)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--check", action="store_true", help="생성 파일이 최신인지 검사만 함")
    parser.add_argument("--test", action="store_true", help="실제 변환·연결 테스트 실행")
    args = parser.parse_args()
    if args.test:
        suite = unittest.defaultTestLoader.discover(str(ROOT / "tests/data"))
        result = unittest.TextTestRunner(verbosity=2).run(suite)
        return 0 if result.wasSuccessful() else 1
    data, pdfs, missing = convert(args.workbook)
    files = output_files(data, pdfs, missing)
    # All parsing and validation finishes before any output file is written.
    stale = []
    for name, content in files.items():
        path = ROOT / name
        if args.check:
            if not path.exists() or path.read_bytes() != content:
                stale.append(name)
        else:
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(content)
    if stale:
        print("재변환이 필요합니다: " + ", ".join(stale))
        return 1
    print(f"{'Checked' if args.check else 'Generated'}: {len(data['lessons'])} lessons, {len(data['spaces'])} spaces, {len(pdfs)} PDFs")
    return 0


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    try:
        sys.exit(main())
    except (ValueError, KeyError, OSError) as error:
        print(f"변환 중단: {error}", file=sys.stderr)
        sys.exit(1)
